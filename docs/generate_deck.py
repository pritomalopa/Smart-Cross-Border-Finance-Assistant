"""Generate 10-slide submission deck. Run: python docs/generate_deck.py"""
from pptx import Presentation
from pptx.util import Inches, Pt
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "FinTechathon_TopicC_Deck.pptx")

SLIDES = [
    ("Smart Cross-Border Finance Assistant",
     "FinTechathon 2026 · International Track · Topic C\nAI as a Financial Participant\nTeam + Universities\nEnglish (official)"),
    ("The Pain (Why C?)",
     "• Intl students: 8-12% remittance fees, opaque FX\n• Missed tuition deadlines, recurring traps (iCloud x3)\n• Zero local credit history, language barriers\n• APEC Shenzhen context: cross-border is the theme"),
    ("Our Solution",
     "Conversational agent: Understand → Plan → Execute\n• Budget + anomaly, FX, 3-route remittance quote\n• YES-confirm transfers, tuition planner, credit guide\n• Trilingual EN/ZH/BN — EN for judges"),
    ("Architecture",
     "frontend/index.html → FastAPI app.py\n→ agent.py (10 intents) + fx_service.py\n→ security.py + store.py (sandbox) + i18n.py\nDemo: demo/run_demo.py → execution_logs.jsonl + audit"),
    ("Live Demo (2 min)",
     "1. Budget summary + anomaly flagged\n2. 100 USD→CNY + Quote 50000 BDT→CNY (CIPS best)\n3. Send→YES→Tx id\nScreenshots + http://127.0.0.1:8000"),
    ("Security & Compliance (30%)",
     "• Tiers: ≤1000 direct / 1000+ YES / >20000 block\n• 12 injection regexes, PII masking, audit trail\n• Demo: 'Ignore previous instructions…' → BLOCKED\n• Reproduce: tests_smoke.py + S1/S2 logs"),
    ("Interaction (30%)",
     "• Clarification: missing amount/recipient → ask\n• Correction: YES state machine, NO cancels\n• Fallback + multilingual BN/ZH demo\n• Dialogue efficiency: quote→transfer in 2 turns"),
    ("Impact",
     "• Saves ~400 CNY per remittance (CIPS vs SWIFT)\n• Prevents over-budget + subscription waste 180 CNY/yr\n• Credit path 620→680 in 6 mo\n• Real student pain, measurable"),
    ("Roadmap",
     "• Live FX feed + slippage\n• e-KYC + sanctions screening\n• On-chain RMB settlement (Topic F bridge)\n• Bank API integration with WeBank sandbox"),
    ("Thank You + Q&A",
     "Repo: github.com/pritomalopa/Smart-Cross-Border-Finance-Assistant\nRun: pip install -r requirements.txt; python app.py\nContact + team photo"),
]

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
for title, body in SLIDES:
    sl = prs.slides.add_slide(prs.slide_layouts[1])
    sl.shapes.title.text = title
    sl.placeholders[1].text = body
prs.save(OUT)
print(f"Saved {OUT} ({len(SLIDES)} slides)")
